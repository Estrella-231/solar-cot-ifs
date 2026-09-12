"""Read-only train/validation source flatline diagnostics; never changes QC masks.

The 60 minute flag describes 12 equal interval-ending 5 minute records.
It is an evidence tag, not an exclusion rule. No test numerical payload is
returned from the workbook readers or parsed from the source cohort.
"""
import argparse
from collections import Counter, defaultdict
import csv
from datetime import datetime, timedelta, timezone
import hashlib
import json
import math
from pathlib import Path

import openpyxl

BJT = timezone(timedelta(hours=8))
START = datetime(2024, 4, 2, tzinfo=BJT)
BOUNDARY = datetime(2025, 7, 1, tzinfo=BJT)
END = datetime(2025, 10, 1, tzinfo=BJT)
SOURCE_SHA = {
    "sili": "c13c6a066cc7ca3ee0d902bb8a9d6c8ea1cc1618659e86eafa4971bbe9a879f4",
    "zhujia": "38e4b708a009a69d16bc49ff80ea330a88a6b32e1d69ff6ee53dc49ee13783a0",
}
WORKBOOK_NAMES = {"sili": "四里光伏.xlsx", "zhujia": "竺家光伏.xlsx"}


def sha(path):
    h = hashlib.sha256()
    with Path(path).open("rb") as f:
        for b in iter(lambda: f.read(1048576), b""):
            h.update(b)
    return h.hexdigest()


def split_of(t):
    assert START <= t < END
    return "train" if t < BOUNDARY else "validation"


def load_source(path, station):
    assert sha(path) == SOURCE_SHA[station], station
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    metadata = []
    source = {}
    stats = Counter()
    for sheet in book:
        header = [c.value for c in next(sheet.iter_rows(max_row=1))]
        identity_columns = header.index("DATA_DATE") + 1
        assert header[1:identity_columns] in (["OBJ_TYPE", "DATA_DATE"], ["DATA_DATE"])
        expected = ["VAL" + f"{m//60:02d}{m%60:02d}" for m in range(5, 1441, 5)]
        assert header[identity_columns:] == expected
        dates = {}
        # This pass requests identity/date cells only, never numerical columns.
        previous = None
        for cells in sheet.iter_rows(min_row=2, max_col=identity_columns):
            date = datetime.strptime(str(cells[-1].value), "%Y-%m-%d").replace(tzinfo=BJT)
            assert previous is None or date >= previous, "Source day rows not sorted"
            previous = date
            if START - timedelta(days=1) <= date < END:
                if "OBJ_TYPE" in header:
                    assert cells[header.index("OBJ_TYPE")].value == "IRRADIANCE"
                dates[cells[-1].row] = date
        assert dates
        first, last = min(dates), max(dates)
        assert dates[first] == START - timedelta(days=1)
        assert dates[last] == END - timedelta(days=1)
        assert set(dates) == set(range(first, last + 1))
        metadata.append({"sheet": sheet.title, "total_rows_including_header": sheet.max_row,
                         "columns": sheet.max_column, "selected_first_row": first,
                         "selected_last_row": last, "selected_day_rows": len(dates),
                         "identity_columns": header[:identity_columns], "date_type": "text YYYY-MM-DD",
                         "measurement_type": "VALHHMM interval end BJT"})

        def accept(cells, col_start):
            source_row = cells[0].row
            date = dates[source_row]
            for col, cell in enumerate(cells, col_start):
                minutes = (col - identity_columns) * 5
                t = date + timedelta(minutes=minutes)
                assert START <= t < END, "Reader exposed outside-scope numerical cell"
                key = split_of(t)
                stats[key + "_raw_records"] += 1
                raw = cell.value
                try:
                    value = float(raw)
                except (ValueError, TypeError):
                    value = None
                if value is None or not math.isfinite(value):
                    value = None
                    stats[key + "_missing_or_nonfinite"] += 1
                elif value < 0:
                    stats[key + "_negative"] += 1
                elif value == 0:
                    stats[key + "_zero"] += 1
                else:
                    stats[key + "_positive"] += 1
                assert t not in source, "Duplicate source timestamp needs separate audit"
                source[t] = {"value": value, "sheet": sheet.title, "cell": cell.coordinate}

        # Previous-day midnight belongs to train; only VAL2400 is requested.
        final_column = identity_columns + 288
        first_column = identity_columns + 1
        accept(next(sheet.iter_rows(min_row=first, max_row=first, min_col=final_column, max_col=final_column)), final_column)
        for cells in sheet.iter_rows(min_row=first + 1, max_row=last - 1, min_col=first_column, max_col=final_column):
            accept(cells, first_column)
        # The last VAL2400 is test 2025-10-01 00:00 and is not requested.
        accept(next(sheet.iter_rows(min_row=last, max_row=last, min_col=first_column, max_col=final_column - 1)), first_column)
    book.close()
    return source, metadata, dict(stats)


def positive_runs(source):
    runs, current = [], []
    for t in sorted(source):
        value = source[t]["value"]
        positive = value is not None and math.isfinite(value) and value > 0
        extends = (current and positive and t == current[-1] + timedelta(minutes=5)
                   and split_of(t) == split_of(current[-1])
                   and value == source[current[-1]]["value"])
        if current and not extends:
            runs.append(current)
            current = []
        if positive:
            current.append(t)
    if current:
        runs.append(current)
    return runs


def load_cohort(path):
    counts, labels = Counter(), {}
    with Path(path).open(newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            if row["split"] == "test":
                continue  # Do not parse any test label or target fields.
            assert row["split"] in ("train", "validation")
            t = datetime.fromisoformat(row["target_time_bjt"])
            assert START <= t < END and row["split"] == split_of(t)
            key = row["station"], t
            value = float(row["observed_ghi_15min_wm2"])
            if key in labels:
                assert labels[key] == value
            labels[key] = value
            counts[key] += 1
    assert sum(counts.values()) == 538053 + 99849
    return counts, labels


def length_bin(n):
    if n == 1: return "1 record"
    if n == 2: return "2 records"
    if n == 3: return "3 records"
    if n <= 6: return "4-6 records"
    if n <= 11: return "7-11 records"
    if n <= 23: return "12-23 records"
    if n <= 47: return "24-47 records"
    return "48+ records"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--sili", type=Path, required=True)
    parser.add_argument("--zhujia", type=Path, required=True)
    parser.add_argument("--cohort", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    counts, labels = load_cohort(args.cohort)
    report = {"audit_time_utc": datetime.now(timezone.utc).isoformat(),
              "status": "DIAGNOSTIC_FLAGS_ONLY_NO_QC_EXCLUSIONS",
              "scope": {"start_bjt_inclusive": START.isoformat(), "train_end_bjt_exclusive": BOUNDARY.isoformat(),
                        "validation_end_bjt_exclusive": END.isoformat(), "test_numeric_values_analyzed": False},
              "contract": {"positive": "parsed finite GHI > 0; exact numeric equality, no rounding/tolerance",
                           "continuous": "consecutive interval-end timestamps differ by exactly5min; split boundary breaks run",
                           "flag": "at least12 consecutive equal positive 5min interval records, total interval support>=60min",
                           "duration_fields": "support_minutes=5*n; endpoint_span_minutes=5*(n-1)",
                           "effect": "any of target-10/-5/0 is in run; same station/target across init and lead counted separately",
                           "not_an_exclusion_rule": True},
              "source_sha256": SOURCE_SHA, "script_sha256": sha(__file__),
              "cohort_sha256": sha(args.cohort), "sources": {}, "groups": [], "flagged_runs": []}
    all_affected = set()
    for station in ("sili", "zhujia"):
        source, sheets, raw_stats = load_source(getattr(args, station), station)
        report["sources"][station] = {"authoritative_path": "/home/Data_Pool_3/chenyi/GHI/data/hunan/raw/stations/" + WORKBOOK_NAMES[station],
                                         "verified_cache_path": "/tmp/hunan_m0_staging_20260903/" + WORKBOOK_NAMES[station],
                                         "verified_local_copy": str(getattr(args, station)), "sheets": sheets,
                                         "cache_matches_registered_authority_sha": True, "authority_file_rehashed_this_audit": False,
                                         "raw_statistics": raw_stats}
        runs = positive_runs(source)
        for split in ("train", "validation"):
            selected = [r for r in runs if split_of(r[0]) == split]
            histogram = Counter(length_bin(len(r)) for r in selected)
            flagged = [r for r in selected if len(r) >= 12]
            affected = set()
            fully_affected = set()
            for r in flagged:
                run_set = set(r)
                targets = {t + timedelta(minutes=(-t.minute) % 15) for t in r}
                keys = {(station, t) for t in targets if (station, t) in counts}
                full = {k for k in keys if all(k[1] - timedelta(minutes=m) in run_set for m in (10, 5, 0))}
                affected |= keys
                fully_affected |= full
                discrepancies = []
                for key in sorted(keys):
                    times = [key[1] - timedelta(minutes=m) for m in (10, 5, 0)]
                    vals = [source.get(t, {}).get("value") for t in times]
                    if all(v is not None and math.isfinite(v) and v >= 0 for v in vals):
                        delta = abs(sum(vals) / 3 - labels[key])
                        if delta > 1e-4:
                            discrepancies.append({"target_time_bjt": key[1].isoformat(), "absolute_difference_wm2": delta})
                    else:
                        discrepancies.append({"target_time_bjt": key[1].isoformat(), "error": "cohort label refers to incomplete source triple"})
                assert not discrepancies, discrepancies
                first, last = r[0], r[-1]
                report["flagged_runs"].append({"station": station, "split": split,
                    "interval_start_bjt_exclusive": (first - timedelta(minutes=5)).isoformat(),
                    "first_interval_end_bjt": first.isoformat(), "last_interval_end_bjt": last.isoformat(),
                    "records": len(r), "support_minutes": 5*len(r), "endpoint_span_minutes": 5*(len(r)-1),
                    "ghi_wm2": source[first]["value"], "first_source_cell": source[first], "last_source_cell": source[last],
                    "potential_15min_targets": len(targets), "affected_cohort_unique_targets": len(keys),
                    "affected_cohort_rows": sum(counts[k] for k in keys),
                    "fully_flat_cohort_unique_targets": len(full), "fully_flat_cohort_rows": sum(counts[k] for k in full),
                    "affected_target_times_bjt": [k[1].isoformat() for k in sorted(keys)], "cohort_aggregation_check": "PASS"})
            all_affected |= affected
            group_keys = {k for k in counts if k[0] == station and split_of(k[1]) == split}
            report["groups"].append({"station": station, "split": split,
                "positive_runs": len(selected), "positive_run_length_histogram": dict(histogram),
                "flagged_runs": len(flagged), "flagged_source_records": sum(map(len, flagged)),
                "flagged_unique_source_dates": len({t.date() for r in flagged for t in r}),
                "max_support_minutes": max([len(r)*5 for r in flagged], default=0),
                "flagged_value_distribution": dict(Counter(str(source[r[0]]["value"]) for r in flagged)),
                "affected_cohort_unique_targets": len(affected), "affected_cohort_rows": sum(counts[k] for k in affected),
                "fully_flat_cohort_unique_targets": len(fully_affected), "fully_flat_cohort_rows": sum(counts[k] for k in fully_affected),
                "all_cohort_unique_targets": len(group_keys), "all_cohort_rows": sum(counts[k] for k in group_keys)})
    report["total"] = {"flagged_runs": len(report["flagged_runs"]),
                       "affected_cohort_unique_targets": len(all_affected),
                       "affected_cohort_rows": sum(counts[k] for k in all_affected),
                       "all_cohort_rows": sum(counts.values()), "labels_modified": False, "cohort_modified": False}
    report["top10_by_support_duration"] = sorted(report["flagged_runs"],
        key=lambda r: (-r["support_minutes"], r["station"], r["first_interval_end_bjt"]))[:10]
    args.output.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(args.output), "groups": report["groups"], "total": report["total"]}, ensure_ascii=True))


if __name__ == "__main__":
    main()
